"""Bounded local extraction. Uploaded content is data, never executable instructions."""
import io,re,zipfile
from pathlib import Path
from fastapi import HTTPException

MAX_BYTES=15*1024*1024


def extract_document(data,name):
    suffix=Path(name or '').suffix.lower()
    if len(data)>MAX_BYTES:raise HTTPException(413,'Use a PDF or PPTX smaller than 15 MB')
    if suffix=='.pdf':
        import pymupdf
        try:
            with pymupdf.open(stream=data,filetype='pdf') as document:
                if document.needs_pass:raise ValueError('Password protected')
                if not 1<=len(document)<=60:raise ValueError('Use 1–60 pages')
                sections=[{'number':i+1,'title':f'Page {i+1}','text':page.get_text()[:6000],'notes':''} for i,page in enumerate(document)]
        except Exception as error:raise HTTPException(422,'Could not read this PDF. Use an unlocked PDF with 1–60 pages.') from error
    elif suffix=='.pptx':
        from pptx import Presentation
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as package:
                entries=package.infolist()
                if len(entries)>5000 or sum(item.file_size for item in entries)>80*1024*1024 or any(item.flag_bits&1 for item in entries):raise ValueError('Archive exceeds limits')
            deck=Presentation(io.BytesIO(data))
            if not 1<=len(deck.slides)<=60:raise ValueError('Use 1–60 slides')
            sections=[]
            for index,slide in enumerate(deck.slides):
                pieces=[]
                for shape in slide.shapes:
                    if shape.has_text_frame:pieces.append(shape.text)
                    if shape.has_table:pieces.extend(' | '.join(cell.text for cell in row.cells) for row in shape.table.rows)
                sections.append({'number':index+1,'title':(slide.shapes.title.text if slide.shapes.title else f'Slide {index+1}')[:200],
                    'text':'\n'.join(pieces)[:6000],'notes':slide.notes_slide.notes_text_frame.text[:2000] if slide.has_notes_slide and slide.notes_slide.notes_text_frame else ''})
        except Exception as error:raise HTTPException(422,'Could not read this PPTX. Export older PPT files as PPTX or PDF; use 1–60 slides.') from error
    else:raise HTTPException(422,'Upload PDF or PPTX. Export older PPT files as PPTX or PDF first.')
    if sum(len(section['text'])+len(section['notes']) for section in sections)>150000:raise HTTPException(422,'This material has too much text. Upload a shorter section.')
    return suffix[1:],sections


def vocabulary(sections):
    ignored=set('about after again also because before between could every first from have into more most other should their there these they this those through using very what when where which while with would your presentation slide page'.split())
    words=[];seen=set()
    for word in re.findall(r"[^\W\d_]+(?:['’][^\W\d_]+)?",' '.join(section['text'] for section in sections),re.UNICODE):
        key=word.casefold()
        if len(word)>=5 and key not in ignored and key not in seen:words.append(word);seen.add(key)
    return words[:24]
